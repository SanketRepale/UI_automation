from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".rst"}:
        return content.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if suffix == ".docx":
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + tables)
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))
        return "\n".join(" | ".join(row) for row in rows)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            lines.extend(" | ".join(str(value or "") for value in row) for row in sheet.iter_rows(values_only=True))
        return "\n".join(lines)
    raise ValueError(f"Unsupported file type: {suffix or 'no extension'}")


def extract_application_details(text: str) -> dict[str, Any]:
    """Extract explicitly stated setup values locally; never send credentials to an LLM."""
    urls = re.findall(r"https?://[^\s<>\]\[()\"']+", text, flags=re.IGNORECASE)
    clean_urls = []
    values: dict[str, str | None] = {"username": None, "password": None}
    for raw_url in urls:
        candidate = raw_url.rstrip(".,;:")
        try:
            parts = urlsplit(candidate)
            if parts.scheme and parts.hostname:
                if parts.username and not values["username"]:
                    values["username"] = parts.username
                if parts.password and not values["password"]:
                    values["password"] = parts.password
                safe_query = [(key, "[REDACTED]" if re.search(r"password|passwd|secret|token|api.?key|credential", key, re.I) else value) for key, value in parse_qsl(parts.query, keep_blank_values=True)]
                clean_urls.append(urlunsplit((parts.scheme, parts.hostname + (f":{parts.port}" if parts.port else ""), parts.path, urlencode(safe_query, doseq=True), parts.fragment)))
        except ValueError:
            continue

    for key, pattern in {
        "username": r"(?im)^\s*(?:username|user\s*name|user\s*id|login\s*id|email)\s*[:=|]\s*([^\s,;|]+)",
        "password": r"(?im)^\s*(?:password|passwd|passcode)\s*[:=|]\s*([^\s,;|]+)",
    }.items():
        match = re.search(pattern, text)
        if match:
            values[key] = match.group(1).strip("\"'`")
    table_rows = _table_rows(text)
    for header_index, cells in enumerate(table_rows[:-1]):
        normalized_headers = [cell.strip().lower().replace("_", " ") for cell in cells]
        row = table_rows[header_index + 1]
        if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row if cell.strip()):
            if header_index + 2 < len(table_rows):
                row = table_rows[header_index + 2]
        if len(row) != len(cells):
            continue
        for key, names in {"username": {"username", "user name", "user id", "login id", "email"}, "password": {"password", "passwd", "passcode"}}.items():
            if not values[key]:
                column = next((index for index, name in enumerate(normalized_headers) if name in names), None)
                if column is not None and row[column].strip():
                    values[key] = row[column].strip().strip("\"'`")

    lowered = text.lower()
    auth_type = None
    if re.search(r"\b(no authentication|unauthenticated|public application)\b", lowered):
        auth_type = "No Authentication"
    elif re.search(r"\b(single sign[ -]?on|\bsso\b|saml|oauth login)", lowered):
        auth_type = "SSO"
    elif values["username"] or values["password"] or re.search(r"username\s*(?:and|&)\s*password", lowered):
        auth_type = "Username & Password"

    environment_match = re.search(r"(?im)^\s*(?:environment|deployment environment)\s*[:=|]\s*(QA|DEV|UAT)\b", text)
    browser_match = re.search(r"(?im)^\s*browser\s*[:=|]\s*(chromium|chrome|firefox|webkit|safari)\b", text)
    mode_match = re.search(r"(?im)^\s*(?:execution\s*mode|browser\s*mode)\s*[:=|]\s*(headed|headless)\b", text)
    username_label = re.search(r"(?im)^\s*(?:username|user)\s+field\s+label\s*[:=|]\s*(.+?)\s*$", text)
    password_label = re.search(r"(?im)^\s*password\s+field\s+label\s*[:=|]\s*(.+?)\s*$", text)
    submit_label = re.search(r"(?im)^\s*(?:sign[ -]?in|login)\s+button\s+label\s*[:=|]\s*(.+?)\s*$", text)
    return {
        "application_url": clean_urls[0] if clean_urls else None,
        "urls_found": list(dict.fromkeys(clean_urls)),
        "authentication_type": auth_type,
        "username": values["username"],
        "password": values["password"],
        "environment": environment_match.group(1).upper() if environment_match else None,
        "browser": {"chrome": "chromium", "safari": "webkit"}.get(browser_match.group(1).lower(), browser_match.group(1).lower()) if browser_match else None,
        "execution_mode": mode_match.group(1).lower() if mode_match else None,
        "username_label": username_label.group(1).strip() if username_label else None,
        "password_label": password_label.group(1).strip() if password_label else None,
        "submit_label": submit_label.group(1).strip() if submit_label else None,
    }


def redact_credentials(text: str) -> str:
    """Remove credential values from document text before constructing any LLM prompt."""
    lines = text.splitlines()
    active_sensitive_columns: tuple[str, list[int]] | None = None
    for line_index, line in enumerate(lines):
        cells = _split_table_row(line)
        if not cells:
            active_sensitive_columns = None
            continue
        separator, row = cells
        normalized = [cell.strip().lower().replace("_", " ") for cell in row]
        secret_headers = {"username", "user name", "user id", "login id", "email", "password", "passwd", "passcode"}
        header_columns = [index for index, cell in enumerate(normalized) if cell in secret_headers]
        if header_columns:
            active_sensitive_columns = (separator, header_columns)
            continue
        if active_sensitive_columns:
            active_separator, columns = active_sensitive_columns
            if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row if cell.strip()):
                continue
            if separator != active_separator or len(row) <= max(columns):
                active_sensitive_columns = None
                continue
            for column in columns:
                row[column] = "[REDACTED]"
            lines[line_index] = f" {separator} ".join(row)
    text = "\n".join(lines)
    sensitive = r"(?:username|user_name|user name|user id|login id|password|passwd|passcode|client_secret|access_token|api[_-]?key|token)"
    text = re.sub(rf"(?i)([\"']?{sensitive}[\"']?\s*[:=]\s*[\"']?)([^\s\"',;]+)", r"\1[REDACTED]", text)
    text = re.sub(rf"(?im)(\b{sensitive}\b\s*[|,]\s*)[^\r\n|,]+", r"\1[REDACTED]", text)
    text = re.sub(r"(?i)(https?://)[^/\s@]+@", r"\1[REDACTED]@", text)
    return text


def _table_rows(text: str) -> list[list[str]]:
    return [parts[1] for line in text.splitlines() if (parts := _split_table_row(line))]


def _split_table_row(line: str) -> tuple[str, list[str]] | None:
    for separator in ("|", "\t", ","):
        if separator in line:
            return separator, line.split(separator)
    return None
